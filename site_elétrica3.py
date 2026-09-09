import math
import streamlit as st
from fpdf import FPDF
from docx import Document
from io import BytesIO

st.set_page_config(page_title="Gestor Eletrico Pro", layout="wide")

# --- SISTEMA DE SEGURANÇA E SENHA ---
if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if not st.session_state.autenticado:
    st.title("⚡ Gestor de Projetos Elétricos Avançado")
    st.markdown("### 🔒 Acesso Restrito")
    senha = st.text_input("Digite a senha para acessar o Gestor:", type="password")
    if st.button("Entrar"):
        if senha == "mudar123":  
            st.session_state.autenticado = True
            st.rerun()
        else:
            st.error("Senha incorreta!")
    st.stop()

# --- INICIALIZAÇÃO DE VARIÁVEIS ---
if "comodos" not in st.session_state:
    st.session_state.comodos = []
if "tues_temporarias" not in st.session_state:
    st.session_state.tues_temporarias = []

def gerar_pdf(cliente, obra, responsavel, registro, comodos, circuitos, padrao, inc_padrao):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(190, 10, "LAUDO TECNICO DE DIMENSIONAMENTO ELETRICO", ln=True, align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(190, 5, "Baseado estritamente nas diretrizes da Norma ABNT NBR 5410", ln=True, align="C")
    pdf.line(10, 27, 200, 27)
    pdf.ln(5)
    
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(190, 6, f"CLIENTE: {cliente.upper() if cliente else 'NAO INFORMADO'}", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(190, 5, f"Endereco da Obra: {obra if obra else 'Nao informado'}", ln=True)
    pdf.cell(190, 5, f"Responsavel Tecnico: {responsavel if responsavel else 'Nao informado'} | Registro: {registro if registro else 'Nao informado'}", ln=True)
    pdf.line(10, 47, 200, 47)
    pdf.ln(5)
    
    if inc_padrao:
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(190, 8, f"PADRAO DE ENTRADA RECOMENDADO: {padrao['tipo']}", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(190, 5, f"Disjuntor Geral do Padrao: {padrao['disjuntor']} A | Cabo do Ramal Principal: {padrao['cabo']} mm2", ln=True)
        pdf.ln(5)
    
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(190, 8, "1. Quadro de Distribuicao de Circuitos Otimizado (QGD)", ln=True)
    for circ in circuitos:
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(190, 6, f"Circuito {circ['numero']} - {circ['nome']} ({circ['tensao']}V)".encode('latin-1', 'ignore').decode('latin-1'), ln=True)
        pdf.set_font("Helvetica", "", 10)
        linha2 = f"  -> Potencia: {circ['potencia']:.0f} VA | Corrente: {circ['corrente']:.2f} A | Queda: {circ['queda_tensao']:.2f}%"
        pdf.cell(190, 5, linha2, ln=True)
        pdf.cell(190, 5, f"  -> Condutor: {circ['bitola']} mm2 | Disjuntor: {circ['disjuntor']} A | Agrupamento: {circ['agrupados']} circ.", ln=True)
        pdf.ln(1)
    return pdf.output()

def gerar_word(cliente, obra, responsavel, registro, circuitos, padrao, inc_padrao):
    doc = Document()
    doc.add_heading("MEMORIAL DESCRITIVO E LAUDO ELETRICO", level=1)
    p = doc.add_paragraph()
    p.add_run(f"Cliente: {cliente}\nEndereço: {obra}\n").bold = True
    p.add_run(f"Responsável Técnico: {responsavel} | Registro: {registro}\n")
    
    if inc_padrao:
        doc.add_heading("1. Dimensionamento do Padrão de Entrada Geral", level=2)
        doc.add_paragraph(f"Tipo de Atendimento: {padrao['tipo']}\nDisjuntor Geral Sugerido: {padrao['disjuntor']} A\nCabo do Ramal de Entrada: {padrao['cabo']} mm²")
    
    doc.add_heading("2. Detalhamento dos Circuitos Terminais (QGD)", level=2)
    for circ in circuitos:
        doc.add_paragraph(f"Circuito {circ['numero']} - {circ['nome']}\n"
                          f"Potência: {circ['potencia']:.0f} VA | Tensão: {circ['tensao']} V\n"
                          f"Condutor Final: {circ['bitola']} mm² | Proteção (Disjuntor): {circ['disjuntor']} A\n"
                          f"Queda de Tensão Calculada: {circ['queda_tensao']:.2f}% | Agrupamento no Eletroduto: {circ['agrupados']} circuito(s)")
    
    target = BytesIO()
    doc.save(target)
    return target.getvalue()

st.markdown("""
    <div style="background-color: #111827; padding: 24px; border-radius: 12px; margin-bottom: 30px; text-align: center; border: 1px solid #1e3a8a;">
        <div style="font-size: 40px; margin-bottom: 5px; text-shadow: 0 0 12px #3b82f6;">⚡</div>
        <h1 style="margin: 0; font-size: 28px; font-weight: 800; letter-spacing: 2.5px; color: #ffffff;">RIBEIRO ELÉTRICA <span style="color: #3b82f6;">PRO</span></h1>
        <div style="width: 60px; height: 3px; background: linear-gradient(90deg, #1e3a8a, #3b82f6); margin: 12px auto; border-radius: 2px;"></div>
        <p style="margin: 0; color: #9ca3af; font-size: 13px; font-weight: 500; letter-spacing: 1px;">SISTEMA INTELIGENTE DE DIMENSIONAMENTO • NBR 5410</p>
    </div>
""", unsafe_allow_html=True)
st.markdown("### 👤 Identificação Profissional e do Cliente")
c_cli1, c_cli2 = st.columns(2)
nome_cliente = c_cli1.text_input("Nome do Cliente", placeholder="Ex: Joao Silva")
endereco_obra = c_cli2.text_input("Endereço da Obra", placeholder="Ex: Rua das Palmeiras, 150")

c_prof1, c_prof2 = st.columns(2)
nome_responsavel = c_prof1.text_input("Responsável Técnico", placeholder="Ex: Eng. Pedro Santos")
registro_tecnico = c_prof2.text_input("Registro Profissional (CREA / CFT)", placeholder="Ex: 506.XXX-SP")

st.markdown("#### 🛠️ Escopo do Serviço Contratado")
inclui_padrao_entrada = st.checkbox("Incluir dimensionamento e fornecimento do Padrão de Entrada (Caixa de Medição)", value=True)

st.markdown("---")
col_cadastro, col_projeto = st.columns([1.1, 1.4], gap="large")

with col_cadastro:
    st.markdown("### 📝 Cadastrar Novo Cômodo")
    nome_c = st.text_input("Nome do Cômodo", placeholder="Ex: Suite Master, Cozinha, Sala")
    tipo_c = st.radio("Classificação Base do Ambiente", ["Área Seca (Quartos, Sala)", "Área Úmida/Molhada (Cozinha, Área de Serviço)"])
    possui_suite = st.checkbox("Este cômodo possui um Banheiro Integrado (Suíte)?", value=False)
    
    c1, c2, c3 = st.columns(3)
    larg_c = c1.number_input("Largura (m)", min_value=0.1, value=3.0, step=0.1)
    comp_c = c2.number_input("Comprimento (m)", min_value=0.1, value=4.0, step=0.1)
    dist_c = c3.number_input("Distância ao QGD (m)", min_value=1.0, value=10.0, step=1.0)
        
    tensao_c = st.selectbox("Tensão da Iluminação e Tomadas Gerais (TUGs)", (127, 220))
    agrup_c = st.number_input("Circuitos Agrupados no Eletroduto", min_value=1, value=2, step=1)
    
    # --- NOVO BANCO DE DADOS DE TUES TÍPICAS ---
    st.markdown("#### 🔌 Cargas Especiais / TUEs do Cômodo")
    
    equip_sugestao = st.selectbox(
        "Selecione o Equipamento Típico:",
        [
            "Nenhum / Selecione...",
            "Chuveiro Elétrico",
            "Ar Condicionado",
            "Lava e Seca",
            "Forno Elétrico",
            "Micro-ondas Grande",
            "Torneira Elétrica",
            "Outro (Digitar Manualmente)"
        ]
    )
    
    # Define as potências e tensões comerciais com base na seleção
    opcoes_potencia = [0]
    tensao_sugerida = 220
    mostrar_manual = False
    
    if equip_sugestao == "Chuveiro Elétrico":
        opcoes_potencia = [4400, 5500, 6800, 7500]
        tensao_sugerida = 220
    elif equip_sugestao == "Ar Condicionado":
        opcoes_potencia = [900, 1200, 1600, 2200] # Equivalentes aproximados em W para 9k, 12k, 18k e 24k BTUs
        tensao_sugerida = 220
    elif equip_sugestao == "Lava e Seca":
        opcoes_potencia = [2000, 2500, 3000]
        tensao_sugerida = 127
    elif equip_sugestao == "Forno Elétrico":
        opcoes_potencia = [1500, 2200, 3500]
        tensao_sugerida = 220
    elif equip_sugestao == "Micro-ondas Grande":
        opcoes_potencia = [1200, 1500, 1800]
        tensao_sugerida = 127
    elif equip_sugestao == "Torneira Elétrica":
        opcoes_potencia = [3500, 4500, 5500]
        tensao_sugerida = 220
    elif equip_sugestao == "Outro (Digitar Manualmente)":
        mostrar_manual = True

    col_tue_final, col_tue_v = st.columns([1.5, 0.7])
    
    with col_tue_final:
        if mostrar_manual:
            nome_tue = st.text_input("Nome Customizado:", placeholder="Ex: Hidromassagem")
            w_tue = st.number_input("Potência Customizada (W):", min_value=0, value=0, step=100)
        elif equip_sugestao != "Nenhum / Selecione...":
            nome_tue = equip_sugestao
            w_tue = st.selectbox("Escolha a Potência Comercial (W):", opcoes_potencia)
        else:
            nome_tue = ""
            w_tue = 0
            
    with col_tue_v:
        if equip_sugestao != "Nenhum / Selecione...":
            idx_v = 0 if tensao_sugerida == 220 else 1
            v_tue = st.selectbox("Tensão da TUE (V)", (220, 127), index=idx_v, key="tensao_tue_select")
        else:
            v_tue = 220

    if st.button("➕ Vincular TUE"):
        if nome_tue and w_tue > 0:
            st.session_state.tues_temporarias.append({"equipamento": nome_tue, "potencia": w_tue, "tensao": v_tue})
            st.toast(f"TUE '{nome_tue}' vinculada.")
            
    if st.session_state.tues_temporarias:
        st.write("**TUEs vinculadas provisoriamente:**")
        tues_para_remover = []
        for idx, t in enumerate(st.session_state.tues_temporarias):
            t_col1, t_col2 = st.columns(2)
            t_col1.caption(f"• {t['equipamento']}: {t['potencia']}W ({t['tensao']}V)")
            if t_col2.button("❌", key=f"del_tue_temp_{idx}"):
                tues_para_remover.append(idx)
        if tues_para_remover:
            for index in sorted(tues_para_remover, reverse=True):
                st.session_state.tues_temporarias.pop(index)
            st.rerun()
            
    st.markdown("---")
    if st.button("💾 SALVAR CÔMODO NO PROJETO", type="primary"):
        if nome_c:
            is_molhada = "Úmida" in tipo_c
            area_calc = larg_c * comp_c
            per_calc = 2 * (larg_c + comp_c)
            va_ilum = 100 if area_calc < 6 else 100 + (math.floor((area_calc - 6) / 4) * 60)
            
            va_tugs_quarto = 0
            va_tugs_banheiro = 0
            
            if possui_suite:
                q_tugs_q = math.ceil(per_calc / 5)
                va_tugs_quarto = q_tugs_q * 100
                va_tugs_banheiro = 600
                is_molhada = False
            else:
                q_tugs = math.ceil(per_calc / 3.5) if is_molhada else math.ceil(per_calc / 5)
                va_tugs_quarto = ((3 * 600) + ((q_tugs - 3) * 100) if q_tugs > 3 else q_tugs * 600) if is_molhada else q_tugs * 100
                
            st.session_state.comodos.append({
                "nome": nome_c, "molhada": is_molhada, "largura": larg_c, "comprimento": comp_c,
                "area": area_calc, "perimetro": per_calc, "tensao": tensao_c, "distancia": dist_c,
                "va_ilum": va_ilum, "va_tugs": va_tugs_quarto, "agrupados": agrup_c,
                "is_suite": possui_suite, "va_banheiro": va_tugs_banheiro,
                "tues": list(st.session_state.tues_temporarias)
            })
            st.session_state.tues_temporarias = []
            st.rerun()

    if st.button("🗑️ Limpar Todo o Projeto"):
        st.session_state.comodos = []
        st.session_state.tues_temporarias = []
        st.rerun()
with col_projeto:
    st.markdown("### 📋 Quadro de Distribuição Otimizado")
    if not st.session_state.comodos:
        st.info("Nenhum cômodo cadastrado.")
    else:
        circuitos = []
        c_num = 1
        
        for v in (127, 220):
            ilum_comodos = [c for c in st.session_state.comodos if c['tensao'] == v]
            if ilum_comodos:
                circuitos.append({
                    "numero": c_num, "nome": f"Iluminacao Geral ({v}V)", "potencia": sum(c['va_ilum'] for c in ilum_comodos),
                    "tensao": v, "tipo": "ILUM", "dr": "RECOMENDADO", "distancia": max(c['distancia'] for c in ilum_comodos), "agrupados": max(c['agrupados'] for c in ilum_comodos)
                })
                c_num += 1
        
        for v in (127, 220):
            umidas = [c for c in st.session_state.comodos if c['tensao'] == v and c['molhada']]
            for u in umidas:
                circuitos.append({
                    "numero": c_num, "nome": f"TUGs Cozinha/Servico - {u['nome']}", "potencia": u['va_tugs'],
                    "tensao": v, "tipo": "TUG", "dr": "OBRIGATORIO", "distancia": u['distancia'], "agrupados": u['agrupados']
                })
                c_num += 1
            
            suites_no_valor = [c for c in st.session_state.comodos if c['tensao'] == v and c['is_suite']]
            for s in suites_no_valor:
                circuitos.append({
                    "numero": c_num, "nome": f"TUG Molhada Banheiro - {s['nome']}", "potencia": s['va_banheiro'],
                    "tensao": v, "tipo": "TUG", "dr": "OBRIGATORIO (Banheiro)", "distancia": s['distancia'], "agrupados": s['agrupados']
                })
                c_num += 1
            
            secas = [c for c in st.session_state.comodos if c['tensao'] == v and not c['molhada'] and not c['is_suite']]
            for s in suites_no_valor:
                secas.append({"nome": f"Quarto {s['nome']}", "va_tugs": s['va_tugs'], "distancia": s['distancia'], "agrupados": s['agrupados']})
                
            c_nome, c_va, c_dist, c_agrup = [], 0, 0, 1
            limite = 1200 if v == 127 else 2500
            
            for s in secas:
                if c_va + s['va_tugs'] <= limite:
                    c_nome.append(s['nome'])
                    c_va += s['va_tugs']
                    c_dist = max(c_dist, s['distancia'])
                    c_agrup = max(c_agrup, s['agrupados'])
                else:
                    if c_va > 0:
                        circuitos.append({"numero": c_num, "nome": f"TUGs Secas Agrupadas ({', '.join(c_nome)})", "potencia": c_va, "tensao": v, "tipo": "TUG", "dr": "RECOMENDADO", "distancia": c_dist, "agrupados": c_agrup})
                        c_num += 1
                    c_nome, c_va, c_dist, c_agrup = [s['nome']], s['va_tugs'], s['distancia'], s['agrupados']
            if c_va > 0:
                circuitos.append({"numero": c_num, "nome": f"TUGs Secas Agrupadas ({', '.join(c_nome)})", "potencia": c_va, "tensao": v, "tipo": "TUG", "dr": "RECOMENDADO", "distancia": c_dist, "agrupados": c_agrup})
                c_num += 1

        for c in st.session_state.comodos:
            for t in c['tues']:
                circuitos.append({
                    "numero": c_num, "nome": f"TUE Exclusiva - {t['equipamento']} ({c['nome']})", "potencia": t['potencia'],
                    "tensao": t['tensao'], "tipo": "TUE", "dr": "OBRIGATORIO", "distancia": c['distancia'], "agrupados": c['agrupados']
                })
                c_num += 1

        pot_total_instalada = sum(circ['potencia'] for circ in circuitos)
        pot_com_demanda_w = pot_total_instalada * 0.5
        padrao = {"tipo": "Monofasico (Ate 12kW)", "disjuntor": 40, "cabo": 10.0}
        if 12000 < pot_com_demanda_w <= 25000:
            padrao = {"tipo": "Bifasico (Ate 25kW)", "disjuntor": 50, "cabo": 16.0}
        elif pot_com_demanda_w > 25000:
            padrao = {"tipo": "Trifasico (Acima de 25kW)", "disjuntor": 63, "cabo": 25.0}

        resumo_disjuntores = {}
        resumo_cabos = {1.5: 0.0, 2.5: 0.0, 4.0: 0.0, 6.0: 0.0, 10.0: 0.0}

        for circ in circuitos:
            circ['corrente'] = circ['potencia'] / circ['tensao']
            f_agrup = 1.0 if circ['agrupados'] == 1 else 0.80 if circ['agrupados'] == 2 else 0.70 if circ['agrupados'] == 3 else 0.65
            b_final = 1.5 if circ['tipo'] == "ILUM" else 2.5
            
            while True:
                cap_base = 17.5 if b_final==1.5 else 24.0 if b_final==2.5 else 32.0 if b_final==4.0 else 41.0 if b_final==6.0 else 57.0
                cap_corrigida = cap_base * f_agrup
                q_v = (2.0 * 0.0178 * circ['distancia'] * circ['corrente']) / b_final
                pct = (q_v / circ['tensao']) * 100.0
                
                if circ['corrente'] <= cap_corrigida and pct <= 4.0:
                    circ['bitola'] = b_final
                    circ['queda_tensao'] = pct
                    break
                else:
                    b_final = 2.5 if b_final==1.5 else 4.0 if b_final==2.5 else 6.0 if b_final==4.0 else 10.0
                    if b_final == 10.0:
                        circ['bitola'] = 10.0
                        circ['queda_tensao'] = pct
                        break

            i_proj = circ['corrente']
            dj_adequado = 10 if i_proj <= 10 else 16 if i_proj <= 16 else 20 if i_proj <= 20 else 25 if i_proj <= 25 else 32 if i_proj <= 32 else 40
            circ['disjuntor'] = dj_adequado
            resumo_disjuntores[dj_adequado] = resumo_disjuntores.get(dj_adequado, 0) + 1
            resumo_cabos[circ['bitola']] += circ['distancia'] * 3.0

        v_aba, d_aba, mat_aba = st.tabs(["🏠 Gerenciar", "🗂️ QGD & Padrao", "📦 Lista de Materiais"])
        
        with v_aba:
            comodo_remover = None
            for idx, co in enumerate(st.session_state.comodos):
                with st.expander(f"📍 {co['nome'].upper()}"):
                    st.write(f"Area: {co['area']:.1f} m2 | Condutores agrupados: {co['agrupados']}")
                    if co.get('is_suite'):
                        st.caption("✅ Configuracao Suite Ativa: Cargas do Banheiro separadas automaticamente das cargas do Quarto.")
                    if st.button("Remover Comodo Completo", key=f"del_comodo_{idx}"):
                        comodo_remover = idx
            if comodo_remover is not None:
                st.session_state.comodos.pop(comodo_remover)
                st.rerun()

        with d_aba:
            if inclui_padrao_entrada:
                st.markdown(f"#### 🏢 Entrada Geral: **{padrao['tipo']}**")
                st.caption(f"Disjuntor Geral da Caixa: **{padrao['disjuntor']}A** | Bitola Geral: **{padrao['cabo']} mm²**")
                st.markdown("---")
            
            for circ in circuitos:
                st.markdown(f"""
                <div style="border:1px solid #ddd; padding:12px; border-radius:6px; margin-bottom:10px; background-color:#1e222b;">
                    <h5 style="margin:0; color:#38bdf8;">Circuito {circ['numero']} - {circ['nome']}</h5>
                    <p style="margin:2px 0; font-size:13px;"><b>Fio Final (Fator Termico):</b> {circ['bitola']} mm² | <b>Disjuntor DIN:</b> {circ['disjuntor']} A</p>
                    <p style="margin:2px 0; font-size:12px; color:#aaa;">Queda de Tensao: {circ['queda_tensao']:.2f}% | DR: {circ['dr']}</p>
                </div>
                """, unsafe_allow_html=True)
                
        with mat_aba:
            st.markdown("#### 🛒 Lista de Compras Estimada")
            if inclui_padrao_entrada:
                st.write(f"• Cabo de Cobre do Padrao ({padrao['cabo']} mm²): **15.0 metros**")
                st.write(f"• Disjuntor Geral do Padrao {padrao['disjuntor']}A: **1 un.**")
            for amp, quant in resumo_disjuntores.items():
                st.write(f"• Disjuntor Termomagnetico DIN {amp}A: **{quant} un.**")
            for bit, metros in resumo_cabos.items():
                if metros > 0:
                    st.write(f"• Cabo Flexivel {bit} mm²: **{metros:.1f} metros**")

        st.markdown("---")
        st.markdown("#### 📂 Exportacao da Documentacao da Obra")
        c_down1, c_down2 = st.columns(2)
        with c_down1:
            dados_pdf = gerar_pdf(nome_cliente, endereco_obra, nome_responsavel, registro_tecnico, st.session_state.comodos, circuitos, padrao, inclui_padrao_entrada)
            st.download_button(label="📥 BAIXAR LAUDO EM PDF", data=bytes(dados_pdf), file_name="laudo_eletrico.pdf", mime="application/pdf")
        with c_down2:
            dados_word = gerar_word(nome_cliente, endereco_obra, nome_responsavel, registro_tecnico, circuitos, padrao, inclui_padrao_entrada)
            st.download_button(label="📝 BAIXAR MEMORIAL EM WORD (.DOCX)", data=dados_word, file_name="memorial_descritivo.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
